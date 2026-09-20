"""Deterministic renderers: PDF, XLSX, DOCX, EML and TXT written from the neutral doc models."""

from __future__ import annotations

import email.policy
import email.utils
import io
import re
import zipfile
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from xml.sax.saxutils import escape

from trueup.simulator.files.doc import (
    Chat,
    Doc,
    Heading,
    KeyValues,
    Para,
    Sheet,
    Table,
    Thread,
)

FIXED_TIME = datetime(2026, 12, 1, 8, 0, tzinfo=UTC)
_ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)
CREATOR = "TrueUp synthetic data"


def render_pdf(doc: Doc, path: Path) -> int:
    """Write the PDF and return its page count."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, TableStyle
    from reportlab.platypus import Table as PdfTable

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=10, leading=14)
    cell = ParagraphStyle("cell", parent=body, fontSize=9, leading=12)
    story: list = [
        Paragraph(escape(doc.title), styles["Title"]),
        Paragraph(escape(doc.subtitle), styles["Heading3"]),
        Spacer(1, 0.15 * inch),
    ]
    for block in doc.blocks:
        if isinstance(block, Heading):
            story.append(Paragraph(escape(block.text), styles["Heading2"]))
        elif isinstance(block, Para):
            story.append(Paragraph(escape(block.text), body))
            story.append(Spacer(1, 0.08 * inch))
        elif isinstance(block, KeyValues):
            data = [[Paragraph(escape(k), cell), Paragraph(escape(v), cell)] for k, v in block.rows]
            grid = PdfTable(data, colWidths=[2.0 * inch, 4.5 * inch])
            grid.setStyle(
                TableStyle(
                    [
                        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.lightgrey),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story += [grid, Spacer(1, 0.12 * inch)]
        elif isinstance(block, Table):
            data = [[Paragraph(f"<b>{escape(c)}</b>", cell) for c in block.columns]]
            data += [[Paragraph(escape(str(v)), cell) for v in row] for row in block.rows]
            grid = PdfTable(data, repeatRows=1)
            grid.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story += [grid, Spacer(1, 0.12 * inch)]
    path.parent.mkdir(parents=True, exist_ok=True)
    template = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        invariant=1,
        title=doc.title,
        author=CREATOR,
        creator=CREATOR,
        subject=doc.subtitle,
    )
    template.build(story)
    from pypdf import PdfReader

    return len(PdfReader(str(path)).pages)


def render_xlsx(sheet: Sheet, path: Path) -> int:
    """Write the workbook and return the number of data rows."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    ws = book.active
    ws.title = sheet.title[:31]
    ws.append(list(sheet.columns))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in sheet.rows:
        ws.append([_cell(v) for v in row])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, Decimal):
                cell.number_format = '"$"#,##0.00'
    for index, column in enumerate(sheet.columns, start=1):
        widest = max([len(column)] + [len(str(r[index - 1])) for r in sheet.rows])
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = min(widest + 2, 60)
    book.properties.creator = CREATOR
    book.properties.lastModifiedBy = CREATOR
    book.properties.created = FIXED_TIME.replace(tzinfo=None)
    book.properties.modified = FIXED_TIME.replace(tzinfo=None)
    book.properties.title = sheet.title
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    _normalize_zip(path)
    return len(sheet.rows)


def _cell(value: object) -> object:
    return value if isinstance(value, Decimal | int) else str(value)


def render_docx(doc: Doc, path: Path) -> int:
    """Write the DOCX and return an estimated page count."""
    from docx import Document

    document = Document()
    document.add_heading(doc.title, level=0)
    document.add_paragraph(doc.subtitle)
    words = len(doc.title.split()) + len(doc.subtitle.split())
    for block in doc.blocks:
        if isinstance(block, Heading):
            document.add_heading(block.text, level=2)
            words += len(block.text.split())
        elif isinstance(block, Para):
            document.add_paragraph(block.text)
            words += len(block.text.split())
        elif isinstance(block, KeyValues):
            table = document.add_table(rows=0, cols=2)
            for key, value in block.rows:
                cells = table.add_row().cells
                cells[0].text, cells[1].text = key, value
                words += len(key.split()) + len(value.split())
        elif isinstance(block, Table):
            table = document.add_table(rows=1, cols=len(block.columns))
            for cell, name in zip(table.rows[0].cells, block.columns, strict=True):
                cell.text = name
            for row in block.rows:
                cells = table.add_row().cells
                for cell, value in zip(cells, row, strict=True):
                    cell.text = str(value)
                    words += len(str(value).split())
    props = document.core_properties
    props.author = props.last_modified_by = CREATOR
    props.title = doc.title
    props.subject = doc.subtitle
    props.created = props.modified = FIXED_TIME.replace(tzinfo=None)
    props.revision = 1
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    _normalize_zip(path)
    return max(1, -(-words // 400))


def render_eml(thread: Thread, path: Path, message_id_seed: str) -> int:
    """Write the thread as one .eml (newest message as headers, earlier ones quoted)."""
    from email.message import EmailMessage

    newest = thread.messages[-1]
    msg = EmailMessage(policy=email.policy.default)
    msg["From"] = newest.sender
    msg["To"] = newest.to
    if newest.cc:
        msg["Cc"] = newest.cc
    msg["Date"] = email.utils.format_datetime(newest.sent)
    msg["Subject"] = newest.subject
    msg["Message-ID"] = f"<{message_id_seed}@northstar.example>"
    if len(thread.messages) > 1:
        msg["In-Reply-To"] = f"<{message_id_seed}-{len(thread.messages) - 1}@northstar.example>"
    quoted = []
    for earlier in reversed(thread.messages[:-1]):
        header = f"On {earlier.sent:%a, %b %d, %Y at %H:%M} UTC, {earlier.sender} wrote:"
        quoted.append(header + "\n" + "\n".join("> " + line for line in earlier.body.splitlines()))
    msg.set_content("\n\n".join([newest.body, *quoted]) + "\n")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(msg.as_bytes(policy=email.policy.default))
    return len(thread.messages)


def render_txt(chat: Chat, path: Path) -> int:
    lines = [f"# {chat.channel}"]
    lines += [f"[{line.sent:%Y-%m-%d %H:%M}] {line.sender}: {line.text}" for line in chat.lines]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return len(chat.lines)


def _pin_modified(data: bytes) -> bytes:
    """openpyxl stamps the save time into core.xml; replace it with the fixed time."""
    fixed = FIXED_TIME.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    return re.sub(
        rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>" + fixed + rb"\g<2>", data
    )


def _normalize_zip(path: Path) -> None:
    """Rewrite an OOXML zip with fixed timestamps so the bytes repeat across runs."""
    source = zipfile.ZipFile(path)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as out:
        for info in source.infolist():
            clean = zipfile.ZipInfo(info.filename, date_time=_ZIP_EPOCH)
            clean.compress_type = zipfile.ZIP_DEFLATED
            clean.external_attr = 0o644 << 16
            out.writestr(clean, _pin_modified(source.read(info.filename)))
    source.close()
    path.write_bytes(buffer.getvalue())
