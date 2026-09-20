"""Read any supported finance file into plain text for the Evidence agent."""

from __future__ import annotations

import email.policy
from decimal import Decimal
from email.parser import BytesParser
from pathlib import Path

SUPPORTED = (".pdf", ".xlsx", ".docx", ".eml", ".txt")


class UnsupportedFile(ValueError):
    """Raised for a file type the readers do not handle."""


def read_text(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".xlsx":
        return _read_xlsx(path)
    if suffix == ".docx":
        return _read_docx(path)
    if suffix == ".eml":
        return _read_eml(path)
    if suffix == ".txt":
        return path.read_text(encoding="utf-8").strip()
    raise UnsupportedFile(f"no reader for {suffix or 'files without an extension'}")


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages).strip()


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in book.worksheets:
        lines.append(f"# {sheet.title}")
        for row in sheet.iter_rows():
            cells = [_format_cell(c.value, c.number_format) for c in row]
            if any(cells):
                lines.append(" | ".join(cells))
    book.close()
    return "\n".join(lines).strip()


def _format_cell(value: object, number_format: str) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int | float):
        amount = Decimal(str(value))
        if "$" in number_format:
            return f"${amount:,.2f}"
        return str(int(amount)) if amount == amount.to_integral_value() else str(amount)
    return str(value)


def _read_docx(path: Path) -> str:
    from docx import Document
    from docx.table import Table

    lines: list[str] = []
    for item in Document(str(path)).iter_inner_content():
        if isinstance(item, Table):
            for row in item.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        elif item.text.strip():
            lines.append(item.text)
    return "\n".join(lines).strip()


def _read_eml(path: Path) -> str:
    message = BytesParser(policy=email.policy.default).parsebytes(path.read_bytes())
    headers = [f"{name}: {message[name]}" for name in ("From", "To", "Cc", "Date", "Subject")]
    part = message.get_body(preferencelist=("plain", "html"))
    body = part.get_content() if part is not None else ""
    return "\n".join([*(h for h in headers if not h.endswith(": None")), "", body.strip()]).strip()
