# ruff: noqa: E501
"""Write the held-out seed to backend/seed_holdout/: python scripts/generate_holdout.py.

Three vendors the pipeline has never seen, with documents written in different wordings and
layouts than the demo vendors. Everything is authored by hand here and the expected outcomes are
worked out by hand (the asserts below are plain arithmetic, not pipeline output). Deterministic:
running it twice produces identical JSON. The hidden answer keys land beside the seed and are read
only by scripts/run_holdout.py.
"""

from __future__ import annotations

import email.utils
import json
import re
import shutil
import sys
import zipfile
from datetime import UTC, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from email.message import EmailMessage  # noqa: E402

from docx import Document  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from pypdf import PdfReader  # noqa: E402
from reportlab.lib import colors  # noqa: E402
from reportlab.lib.pagesizes import letter  # noqa: E402
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle  # noqa: E402

from trueup.ingest.manifest import CaseEntry, FileEntry, FileUniverse  # noqa: E402
from trueup.ingest.readers import read_text  # noqa: E402
from trueup.simulator.files.models import RelevanceEntry, RelevanceTruth  # noqa: E402
from trueup.simulator.scenario_models import (  # noqa: E402
    HistoricalTruth,
    OutreachResponse,
    ScenarioEvent,
    StaticCompanyData,
)

OUT = ROOT / "seed_holdout"
FIXED = datetime(2026, 12, 1, 8, 0, tzinfo=UTC)
D = Decimal

# ---- the three cases, in plain words --------------------------------------------------------
#
# Terrastack Compute: GPU compute billed per compute-hour. USD 2.50 an hour, and from 1 Oct 2026 a
#   contractual 8 percent uplift to 2.70. Accruals in Oct and Nov were booked at 2.50 and missed.
#   Hours: Sep 5,800, Oct 6,200, Nov 6,500. December's console export stops on the 21st (4,690
#   hours); the owner later confirms 6,900 hours for the month. December invoice: 6,900 x 2.70.
#
# Larkspur Design Studio: a fixed-fee statement of work paid per accepted deliverable, not to
#   exceed USD 96,000. D1 18,000 (Oct), D2 22,500 (Nov, but the Nov invoice adds 1,200 of
#   unapproved travel), D3 27,450 accepted in December, D4 28,050 later. D3 is above the 25,000
#   Controller limit.
#
# Corvid Security: a monthly subscription. Fee 3,250; an emailed amendment sets 3,900 from
#   16 Dec 2026 with the transition month prorated by calendar day: 15 days at 3,250 plus 16 days
#   at 3,900 over 31 days.

HOURS = {"2026-09": 5800, "2026-10": 6200, "2026-11": 6500, "2026-12": 6900}
PARTIAL_DEC_HOURS = 4690
BASE, UPLIFTED = D("2.50"), D("2.70")
assert BASE * D("1.08") == UPLIFTED
TS_INVOICE = {
    "2026-09": D("14500.00"),
    "2026-10": D("16740.00"),
    "2026-11": D("17550.00"),
    "2026-12": D("18630.00"),
}
TS_ACCRUAL = {"2026-09": D("14500.00"), "2026-10": D("15500.00"), "2026-11": D("16250.00")}
assert TS_INVOICE["2026-09"] == HOURS["2026-09"] * BASE
assert TS_INVOICE["2026-10"] == HOURS["2026-10"] * UPLIFTED
assert TS_INVOICE["2026-11"] == HOURS["2026-11"] * UPLIFTED
assert TS_INVOICE["2026-12"] == HOURS["2026-12"] * UPLIFTED
assert all(TS_ACCRUAL[p] == HOURS[p] * BASE for p in TS_ACCRUAL)

LS_D = {"D1": D("18000.00"), "D2": D("22500.00"), "D3": D("27450.00"), "D4": D("28050.00")}
LS_TOTAL = D("96000.00")
assert sum(LS_D.values()) == LS_TOTAL
LS_NOV_EXPENSES = D("1200.00")
LS_NOV_INVOICE = LS_D["D2"] + LS_NOV_EXPENSES

CV_OLD, CV_NEW = D("3250.00"), D("3900.00")
CV_DEC_RAW = CV_OLD * 15 / 31 + CV_NEW * 16 / 31
CV_DEC = CV_DEC_RAW.quantize(D("0.01"))
assert CV_DEC == D("3585.48")
CV_FIRST, CV_SECOND = (CV_OLD * 15 / 31).quantize(D("0.01")), (CV_NEW * 16 / 31).quantize(D("0.01"))
assert CV_FIRST + CV_SECOND == CV_DEC


def ts(text: str) -> datetime:
    return datetime.fromisoformat(text.replace("Z", "+00:00"))


# ---- document builders ----------------------------------------------------------------------


def _esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def make_pdf(path: Path, title: str, blocks: list[tuple[str, Any]]) -> None:
    styles = getSampleStyleSheet()
    small = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=9, leading=11)
    doc = SimpleDocTemplate(
        str(path), pagesize=letter, title=title, author="holdout", invariant=1,
        leftMargin=54, rightMargin=54, topMargin=54, bottomMargin=54,
    )  # fmt: skip
    story: list[Any] = [Paragraph(_esc(title), styles["Title"])]
    for kind, content in blocks:
        if kind == "h":
            story.append(Paragraph(_esc(content), styles["Heading3"]))
        elif kind == "p":
            story.append(Paragraph(_esc(content), styles["BodyText"]))
        elif kind == "table":
            rows = [[Paragraph(_esc(str(c)), small) for c in row] for row in content]
            table = Table(rows, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.whitesmoke),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            story.append(table)
            story.append(Spacer(1, 8))
    doc.build(story)


def make_docx(path: Path, title: str, blocks: list[tuple[str, Any]]) -> None:
    doc = Document()
    doc.core_properties.author = "holdout"
    doc.core_properties.created = FIXED
    doc.core_properties.modified = FIXED
    doc.core_properties.last_modified_by = "holdout"
    doc.add_heading(title, level=1)
    for kind, content in blocks:
        if kind == "h":
            doc.add_heading(content, level=2)
        elif kind == "p":
            doc.add_paragraph(content)
        elif kind == "table":
            table = doc.add_table(rows=len(content), cols=len(content[0]))
            table.style = "Table Grid"
            for r, row in enumerate(content):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = str(cell)
    doc.save(str(path))
    _fix_zip_stamps(path)


def make_xlsx(path: Path, sheets: dict[str, list[list[Any]]]) -> None:
    book = Workbook()
    book.remove(book.active)
    book.properties.creator = "holdout"
    book.properties.created = FIXED.replace(tzinfo=None)
    book.properties.modified = FIXED.replace(tzinfo=None)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        for row in rows:
            sheet.append(row)
    book.save(str(path))
    _fix_zip_stamps(path)


def _fix_zip_stamps(path: Path) -> None:
    """Office writers stamp zip members and core properties with the current time; pin them."""
    with zipfile.ZipFile(path) as source:
        members = [(info.filename, source.read(info.filename)) for info in source.infolist()]
    stamp = b"2026-12-01T08:00:00Z"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for filename, data in members:
            if filename == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:)",
                    rb"\g<1>" + stamp + rb"\g<2>",
                    data,
                )
            info = zipfile.ZipInfo(filename, date_time=(2026, 12, 1, 8, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, data)


def make_eml(path: Path, file_id: str, thread: list[dict[str, str]]) -> None:
    """`thread` is newest first: each item has frm, to, at (ISO), subject, body."""
    top = thread[0]
    body = top["body"].strip()
    for older in thread[1:]:
        when = ts(older["at"]).strftime("%a, %b %d, %Y at %H:%M UTC")
        quoted = "\n".join("> " + line for line in older["body"].strip().splitlines())
        body += f"\n\nOn {when}, {older['frm']} wrote:\n{quoted}"
    msg = EmailMessage()
    msg["From"] = top["frm"]
    msg["To"] = top["to"]
    msg["Date"] = email.utils.format_datetime(ts(top["at"]))
    msg["Subject"] = top["subject"]
    msg["Message-ID"] = f"<{file_id}@northstar.example>"
    msg.set_content(body)
    path.write_bytes(msg.as_bytes())


def make_txt(path: Path, channel: str, lines: list[tuple[str, str, str]]) -> None:
    rows = [f"# #{channel}"] + [f"[{when}] {who}: {text}" for when, who, text in lines]
    path.write_text("\n".join(rows) + "\n", encoding="utf-8")


def money(value: Decimal) -> str:
    return f"{value:,.2f}"


# ---- the documents, per case -------------------------------------------------------------------

Doc = dict[str, Any]


def vendor_master(vendor_id: str, legal: str, contact: str, terms: str, address: str) -> list:
    return [
        ("h", "Supplier record"),
        (
            "table",
            [
                ["Field", "Value"],
                ["Supplier ID", vendor_id],
                ["Legal name", legal],
                ["Status", "Active"],
                ["Payment terms", terms],
                ["Remit-to", address],
                ["Billing contact", contact],
            ],
        ),
    ]


def terrastack_docs() -> list[Doc]:
    console_hours = [
        212, 216, 238, 246, 241, 190, 183,
        209, 221, 247, 253, 242, 195, 188,
        217, 228, 249, 255, 245, 200, 215,
    ]  # fmt: skip
    assert sum(console_hours) == PARTIAL_DEC_HOURS
    usage_rows: list[list[Any]] = [
        ["Terrastack console - daily GPU compute-hours (export generated 2026-12-23)"],
        ["Account: Northstar Analytics, Inc.   Region: us-east   Tier: on-demand"],
        [],
        ["Date", "GPU compute-hours"],
    ]
    for day, hours in enumerate(console_hours, start=1):
        usage_rows.append([f"2026-12-{day:02d}", hours])
    usage_rows.append(["Total, 2026-12-01 to 2026-12-21", PARTIAL_DEC_HOURS])
    usage_rows.append(["Note: export covers the first 21 days only; the month is not complete."])

    gl_rows: list[list[Any]] = [
        ["Posting date", "Journal", "Vendor", "Description", "Debit", "Credit"],
        ["2026-10-01", "GL-TERRASTACK-2026-09", "Terrastack Compute", "GPU compute Sep", 14500, ""],
        ["2026-10-01", "GL-HOLLIS-2026-09", "Hollis Office Supply", "Toner and paper", 812.4, ""],
        ["2026-10-02", "GL-DUNMORE-2026-09", "Dunmore Travel", "Offsite flights", 3390.15, ""],
        ["2026-10-31", "GL-TERRASTACK-2026-10-ACCRUAL", "Terrastack Compute", "Accrual Oct", 15500, ""],
        ["2026-10-31", "GL-RIDGEWAY-2026-10-ACCRUAL", "Ridgeway Legal", "Accrual Oct", 6200, ""],
        ["2026-11-01", "GL-TERRASTACK-2026-10", "Terrastack Compute", "GPU compute Oct", 16740, ""],
        ["2026-11-01", "GL-PINECONE-2026-10", "Pinecone Facilities", "Cleaning, Oct", 2150, ""],
        ["2026-11-03", "GL-HOLLIS-2026-10", "Hollis Office Supply", "Chairs", 4675.9, ""],
        ["2026-11-15", "GL-DUNMORE-2026-11", "Dunmore Travel", "Conference travel", 5820.5, ""],
        ["2026-11-30", "GL-TERRASTACK-2026-11-ACCRUAL", "Terrastack Compute", "Accrual Nov", 16250, ""],
        ["2026-11-30", "GL-RIDGEWAY-2026-11-ACCRUAL", "Ridgeway Legal", "Accrual Nov", 7300, ""],
        ["2026-12-01", "GL-TERRASTACK-2026-11", "Terrastack Compute", "GPU compute Nov", 17550, ""],
        ["2026-12-01", "GL-PINECONE-2026-11", "Pinecone Facilities", "Cleaning, Nov", 2150, ""],
        ["2026-12-02", "GL-RIDGEWAY-2026-11", "Ridgeway Legal", "Legal fees Nov", 7415, ""],
        ["2026-12-03", "GL-HOLLIS-2026-11", "Hollis Office Supply", "Monitors", 2980, ""],
    ]  # fmt: skip

    return [
        {
            "n": "01", "name": "terrastack_msa_signed.pdf", "kind": "agreement", "fmt": "PDF",
            "at": "2026-01-05T09:00:00Z", "role": "SUPPORTS_AMOUNT", "title": "Master Services Agreement",
            "reason": "The signed agreement sets USD 2.50 per compute-hour and the 8 percent uplift to 2.70 from 1 October 2026.",
            "build": lambda p: make_pdf(
                p, "Master Services Agreement - GPU Compute",
                [
                    ("p", "Between Terrastack Compute, Inc. (Terrastack) and Northstar Analytics, Inc. (Customer). Effective 1 January 2026."),
                    ("h", "1. Services"),
                    ("p", "Terrastack provides on-demand GPU compute capacity through its cloud console. Capacity is metered by the hour."),
                    ("h", "6. Charges and rate adjustment"),
                    ("p", "6.1 Customer pays USD 2.50 for each GPU compute-hour (a Compute Hour) consumed, as recorded by the Terrastack metering system."),
                    ("p", "6.2 With effect from 1 October 2026 the charge in Section 6.1 rises by eight percent (8%), from USD 2.50 to USD 2.70 per Compute Hour."),
                    ("p", "6.3 Terrastack invoices monthly in arrears for the Compute Hours consumed in the previous calendar month."),
                    ("h", "12. Term"),
                    ("p", "This Agreement runs from 1 January 2026 to 31 December 2027."),
                    ("p", "Signed for Terrastack: M. Okafor.   Signed for Customer: A. Rivera."),
                ],
            ),
        },
        {
            "n": "02", "name": "gpu_hours_daily_dec_partial.xlsx", "kind": "usage", "fmt": "XLSX",
            "at": "2026-12-23T18:00:00Z", "role": "SUPPORTS_DISCREPANCY", "title": "Daily GPU hours",
            "reason": "The only December usage data, and it stops on the 21st, so the month is incomplete.",
            "build": lambda p: make_xlsx(p, {"Console export": usage_rows}),
        },
        {
            "n": "03", "name": "re_december_gpu_hours_missing.eml", "kind": "email", "fmt": "EML",
            "at": "2026-12-31T08:30:00Z", "role": "SUPPORTS_DISCREPANCY", "title": "December GPU hours",
            "reason": "The owner says the full-month total is not available yet, so the accrual cannot be finalised.",
            "build": lambda p: make_eml(p, "FILE-TERRASTACK-03", [
                {"frm": "Jordan Patel <jordan.patel@northstar.example>", "to": "Riley Kim <riley.kim@northstar.example>",
                 "at": "2026-12-31T08:30:00Z", "subject": "Re: December GPU hours",
                 "body": "Riley - thanks. The console export in the shared folder stops on the 21st.\nDo you have the total compute-hours for the whole of December? I cannot book the December accrual until I have it.\n\nJordan"},
                {"frm": "Riley Kim <riley.kim@northstar.example>", "to": "Jordan Patel <jordan.patel@northstar.example>",
                 "at": "2026-12-28T15:10:00Z", "subject": "Re: December GPU hours",
                 "body": "The console shows what it shows. Terrastack's billing rollup only runs after month end, so I do not have a full-month number today."},
            ]),
        },
        {
            "n": "04", "name": "q4_accrual_review_terrastack.docx", "kind": "prior", "fmt": "DOCX",
            "at": "2026-12-03T09:00:00Z", "role": "SUPPORTS_DISCREPANCY", "title": "Q4 accrual review",
            "reason": "Records that the October and November accruals used 2.50 while the vendor billed 2.70 after the uplift.",
            "build": lambda p: make_docx(
                p, "Q4 accrual review - Terrastack Compute",
                [
                    ("p", "Prepared by Finance Operations on 3 December 2026 for the Controller."),
                    ("h", "What happened"),
                    ("p", "September was accrued and invoiced at USD 2.50 per compute-hour and matched."),
                    ("p", "October: accrued USD 15,500.00 (6,200 hours at 2.50). Invoice TSK-90455 was USD 16,740.00 (6,200 hours at 2.70)."),
                    ("p", "November: accrued USD 16,250.00 (6,500 hours at 2.50). Invoice TSK-90512 was USD 17,550.00 (6,500 hours at 2.70)."),
                    ("h", "Cause"),
                    ("p", "The accruals used the base rate. The 8 percent contractual uplift that took effect on 1 October was not applied."),
                ],
            ),
        },
        {
            "n": "05", "name": "invoice_tsk-90388_sep.pdf", "kind": "invoice", "fmt": "PDF",
            "at": "2026-12-01T08:00:00Z", "role": "HARD_NEGATIVE", "title": "Invoice TSK-90388",
            "reason": "A September invoice at the pre-uplift rate covers a closed period and would mislead the December estimate.",
            "build": lambda p: make_pdf(
                p, "Invoice TSK-90388",
                [
                    ("p", "Terrastack Compute, Inc.   Bill to: Northstar Analytics, Inc.   Invoice date: 30 September 2026"),
                    ("table", [["Description", "Period", "Hours", "Rate (USD)", "Amount (USD)"],
                               ["GPU compute-hours", "1 Sep - 30 Sep 2026", "5,800", "2.50", "14,500.00"]]),
                    ("p", "Total due: USD 14,500.00. Terms: net 30."),
                ],
            ),
        },
        {
            "n": "06", "name": "terrastack_order_form_2025.pdf", "kind": "agreement", "fmt": "PDF",
            "at": "2026-12-01T08:00:00Z", "role": "HARD_NEGATIVE", "title": "Order Form OF-2025-114",
            "reason": "A 2025 order form at USD 2.20 per hour that the 2026 master agreement replaced.",
            "build": lambda p: make_pdf(
                p, "Order Form OF-2025-114",
                [
                    ("p", "Terrastack Compute, Inc. and Northstar Analytics, Inc. Term: 1 July 2025 to 31 December 2025."),
                    ("p", "GPU compute-hours are charged at USD 2.20 per hour on the on-demand tier."),
                    ("p", "This order form is replaced by the Master Services Agreement dated 1 January 2026 and has no effect after 31 December 2025."),
                ],
            ),
        },
        {
            "n": "07", "name": "terrastack_storage_invoice_dec.pdf", "kind": "invoice", "fmt": "PDF",
            "at": "2026-12-20T10:00:00Z", "role": "HARD_NEGATIVE", "title": "Invoice TSL-2210",
            "reason": "An invoice from a different vendor with a similar name, for object storage.",
            "build": lambda p: make_pdf(
                p, "Invoice TSL-2210",
                [
                    ("p", "Terrastack Storage Ltd.   Bill to: Northstar Analytics, Inc.   Invoice date: 20 December 2026"),
                    ("table", [["Description", "Period", "Quantity", "Rate (USD)", "Amount (USD)"],
                               ["Object storage, TB-months", "1 Nov - 30 Nov 2026", "41.2", "100.00", "4,120.00"]]),
                    ("p", "Total due: USD 4,120.00."),
                ],
            ),
        },
        {
            "n": "08", "name": "vendor_master_terrastack.pdf", "kind": "vendor", "fmt": "PDF",
            "at": "2026-12-01T08:00:00Z", "role": "NOISE", "title": "Supplier record",
            "reason": "Static supplier identity and payment terms carry no amount or period evidence.",
            "build": lambda p: make_pdf(
                p, "Supplier master - Terrastack Compute",
                vendor_master("VEN-TERRASTACK", "Terrastack Compute, Inc.", "billing@terrastack.example",
                              "Net 30", "410 Harbor Way, Seattle, WA 98101"),
            ),
        },
        {
            "n": "09", "name": "slack_export_ml_infra.txt", "kind": "slack", "fmt": "TXT",
            "at": "2026-12-01T08:00:00Z", "role": "NOISE", "title": "Slack export",
            "reason": "Infrastructure chatter about pools and freezes with no accounting facts.",
            "build": lambda p: make_txt(p, "ml-infra", [
                ("2026-12-09 10:12", "ana", "a100 pool is at 80 percent today, moving the nightly eval to the small pool"),
                ("2026-12-09 10:20", "riley", "fine by me, just keep the checkpoints on the shared volume"),
                ("2026-12-14 16:45", "sam", "holiday change freeze starts the 23rd, get your deploys in before then"),
                ("2026-12-15 09:03", "ana", "who owns the terrastack console login for the contractors?"),
            ]),
        },
        {
            "n": "10", "name": "general_ledger_export_dec.xlsx", "kind": "gl", "fmt": "XLSX",
            "at": "2026-12-01T08:00:00Z", "role": "HARD_NEGATIVE", "title": "General ledger",
            "reason": "A bulk export whose rows are mostly other vendors and show only past periods.",
            "build": lambda p: make_xlsx(p, {"GL export": gl_rows}),
        },
        {
            "n": "L1", "name": "invoice_tsk-90671_dec.pdf", "kind": "invoice", "fmt": "PDF",
            "at": "2027-01-06T09:00:00Z", "role": "LATE_ARRIVAL", "title": "Invoice TSK-90671",
            "reason": "The real December invoice arrives after the close and grades the accrual.",
            "build": lambda p: make_pdf(
                p, "Invoice TSK-90671",
                [
                    ("p", "Terrastack Compute, Inc.   Bill to: Northstar Analytics, Inc.   Invoice date: 5 January 2027"),
                    ("table", [["Description", "Period", "Hours", "Rate (USD)", "Amount (USD)"],
                               ["GPU compute-hours", "1 Dec - 31 Dec 2026", "6,900", "2.70", "18,630.00"]]),
                    ("p", "Total due: USD 18,630.00. Terms: net 30."),
                ],
            ),
        },
        {
            "n": "L2", "name": "reply_terrastack_20270103.eml", "kind": "email", "fmt": "EML",
            "at": "2027-01-03T10:00:00Z", "role": "LATE_ARRIVAL", "title": "December GPU hours",
            "reason": "The owner's reply with the full-month hours arrives after the close.",
            "build": lambda p: make_eml(p, "FILE-TERRASTACK-L2", [
                {"frm": "Riley Kim <riley.kim@northstar.example>", "to": "Jordan Patel <jordan.patel@northstar.example>",
                 "at": "2027-01-03T10:00:00Z", "subject": "Re: Terrastack: information needed for month-end",
                 "body": "Jordan - the rollup finished. December came to 6,900 GPU compute-hours for 1 to 31 December.\n\nRiley"},
            ]),
        },
    ]  # fmt: skip


def larkspur_docs() -> list[Doc]:
    fee_table = [
        ["Deliverable", "Description", "Fee (USD)"],
        ["D1", "Discovery and brand audit", money(LS_D["D1"])],
        ["D2", "Brand strategy and naming", money(LS_D["D2"])],
        ["D3", "Visual identity system and usage guidelines", money(LS_D["D3"])],
        ["D4", "Launch collateral and templates", money(LS_D["D4"])],
        ["Total, not to exceed", "", money(LS_TOTAL)],
    ]
    draft_table = [row[:] for row in fee_table]
    draft_table[3][2] = "24,000.00"
    draft_table[5][2] = "92,550.00"
    nov_lines = [
        ["Description", "Amount (USD)"],
        ["Deliverable D2 - Brand strategy and naming", money(LS_D["D2"])],
        ["Travel and materials (reimbursable)", money(LS_NOV_EXPENSES)],
    ]
    ap_rows = [
        ["Larkspur Design Studio - invoice history"],
        ["Invoice", "Date", "Description", "Amount (USD)"],
        ["LDS-0231", "2026-10-31", "Deliverable D1", 18000.00],
        ["LDS-0247", "2026-11-30", "Deliverable D2 plus expenses", 23700.00],
    ]
    return [
        {
            "n": "01", "name": "sow_no2_signed.pdf", "kind": "agreement", "fmt": "PDF",
            "at": "2026-10-02T09:00:00Z", "role": "SUPPORTS_AMOUNT", "title": "Statement of Work No. 2",
            "reason": "The signed statement of work fixes the fee per deliverable and the USD 96,000 ceiling.",
            "build": lambda p: make_pdf(
                p, "Statement of Work No. 2 - Brand System Design",
                [
                    ("p", "Larkspur Design Studio LLC (Studio) for Northstar Analytics, Inc. (Client). Effective 1 October 2026."),
                    ("h", "3. Deliverables and fees"),
                    ("table", fee_table),
                    ("h", "4. Acceptance and expenses"),
                    ("p", "A fee becomes payable only when Client signs the Acceptance Certificate for that Deliverable."),
                    ("p", "Expenses are not reimbursable unless Client approves them in writing before they are incurred."),
                    ("p", "Total fees under this Statement of Work shall not exceed USD 96,000.00."),
                ],
            ),
        },
        {
            "n": "02", "name": "acceptance_certificate_d3.pdf", "kind": "delivery", "fmt": "PDF",
            "at": "2026-12-29T16:30:00Z", "role": "SUPPORTS_AMOUNT", "title": "Acceptance Certificate D3",
            "reason": "Signed proof that D3 was delivered and accepted, with the fee payable.",
            "build": lambda p: make_pdf(
                p, "Acceptance Certificate - Deliverable D3",
                [
                    ("p", "Statement of Work No. 2, Larkspur Design Studio LLC."),
                    ("table", [["Item", "Detail"],
                               ["Deliverable", "D3 - Visual identity system and usage guidelines"],
                               ["Delivered", "14 December 2026"],
                               ["Accepted by", "Priya Nair, Marketing, on 29 December 2026"],
                               ["Fee payable on acceptance", "USD 27,450.00"],
                               ["Expenses approved", "None"]]),
                ],
            ),
        },
        {
            "n": "03", "name": "d3_signoff_thread.eml", "kind": "email", "fmt": "EML",
            "at": "2026-12-30T10:00:00Z", "role": "SUPPORTS_DISCREPANCY", "title": "D3 sign-off",
            "reason": "Confirms D3 acceptance and warns that the studio billed unapproved expenses on D2.",
            "build": lambda p: make_eml(p, "FILE-LARKSPUR-03", [
                {"frm": "Priya Nair <priya.nair@northstar.example>", "to": "Jordan Patel <jordan.patel@northstar.example>",
                 "at": "2026-12-30T10:00:00Z", "subject": "D3 signed off - accrue the fee only",
                 "body": "Jordan - I signed the acceptance certificate for D3 yesterday, it is in the folder.\nHeads up: on D2 they added USD 1,200 of travel and materials to the invoice even though we never approved expenses.\nPlease accrue only the SOW fee for D3 and do not let expenses through.\n\nPriya"},
            ]),
        },
        {
            "n": "04", "name": "november_close_note_larkspur.docx", "kind": "prior", "fmt": "DOCX",
            "at": "2026-12-04T09:00:00Z", "role": "SUPPORTS_DISCREPANCY", "title": "November close note",
            "reason": "Explains that the November invoice exceeded the accepted D2 fee by unapproved expenses.",
            "build": lambda p: make_docx(
                p, "November close note - Larkspur Design Studio",
                [
                    ("p", "Finance Operations, 4 December 2026."),
                    ("p", "D2 was accrued at the accepted fee of USD 22,500.00."),
                    ("p", "Invoice LDS-0247 arrived at USD 23,700.00. The extra USD 1,200.00 is travel and materials that the statement of work does not allow."),
                    ("p", "The difference is held while Marketing reviews it with the studio."),
                ],
            ),
        },
        {
            "n": "05", "name": "sow_no2_draft_v1.docx", "kind": "agreement", "fmt": "DOCX",
            "at": "2026-09-18T09:00:00Z", "role": "HARD_NEGATIVE", "title": "SOW No. 2 draft v1",
            "reason": "An unsigned draft with D3 at USD 24,000 that the signed statement of work replaced.",
            "build": lambda p: make_docx(
                p, "DRAFT v1 - Statement of Work No. 2 (not signed)",
                [
                    ("p", "For discussion only. Superseded by the signed Statement of Work effective 1 October 2026."),
                    ("table", draft_table),
                ],
            ),
        },
        {
            "n": "06", "name": "lark_and_spur_studios_invoice_1190.pdf", "kind": "invoice", "fmt": "PDF",
            "at": "2026-12-12T10:00:00Z", "role": "HARD_NEGATIVE", "title": "Invoice 1190",
            "reason": "An invoice from a different vendor with a similar name, for signage.",
            "build": lambda p: make_pdf(
                p, "Invoice 1190",
                [
                    ("p", "Lark & Spur Studios Inc.   Bill to: Northstar Analytics, Inc.   Date: 12 December 2026"),
                    ("table", [["Description", "Amount (USD)"], ["Lobby signage design and install", "9,800.00"]]),
                    ("p", "Total due: USD 9,800.00."),
                ],
            ),
        },
        {
            "n": "07", "name": "invoice_lds-0247_nov.pdf", "kind": "invoice", "fmt": "PDF",
            "at": "2026-12-01T08:00:00Z", "role": "HARD_NEGATIVE", "title": "Invoice LDS-0247",
            "reason": "The November invoice covers a closed period and includes charges the statement of work does not allow.",
            "build": lambda p: make_pdf(
                p, "Invoice LDS-0247",
                [
                    ("p", "Larkspur Design Studio LLC.   Bill to: Northstar Analytics, Inc.   Date: 30 November 2026"),
                    ("table", nov_lines),
                    ("p", f"Total due: USD {money(LS_NOV_INVOICE)}."),
                ],
            ),
        },
        {
            "n": "08", "name": "vendor_master_larkspur.pdf", "kind": "vendor", "fmt": "PDF",
            "at": "2026-12-01T08:00:00Z", "role": "NOISE", "title": "Supplier record",
            "reason": "Static supplier identity and payment terms carry no amount or period evidence.",
            "build": lambda p: make_pdf(
                p, "Supplier master - Larkspur Design Studio",
                vendor_master("VEN-LARKSPUR", "Larkspur Design Studio LLC", "billing@larkspur.example",
                              "Net 30", "88 Mercer Street, New York, NY 10012"),
            ),
        },
        {
            "n": "09", "name": "slack_export_brand_launch.txt", "kind": "slack", "fmt": "TXT",
            "at": "2026-12-01T08:00:00Z", "role": "NOISE", "title": "Slack export",
            "reason": "Creative chatter about colours and fonts with no accounting facts.",
            "build": lambda p: make_txt(p, "brand-launch", [
                ("2026-12-10 11:20", "priya", "the forest green reads too dark on the deck, can we try the lighter one"),
                ("2026-12-10 11:42", "lee", "lighter one looks better on the site too"),
                ("2026-12-16 15:05", "priya", "guidelines PDF is a great read, sending to the whole team"),
            ]),
        },
        {
            "n": "10", "name": "ap_history_larkspur.xlsx", "kind": "ap", "fmt": "XLSX",
            "at": "2026-12-01T08:00:00Z", "role": "HARD_NEGATIVE", "title": "Invoice history",
            "reason": "Past invoices are for closed periods and tempt an estimate from history instead of the acceptance.",
            "build": lambda p: make_xlsx(p, {"Invoice history": ap_rows}),
        },
        {
            "n": "L1", "name": "invoice_lds-0298_d3.pdf", "kind": "invoice", "fmt": "PDF",
            "at": "2027-01-09T09:00:00Z", "role": "LATE_ARRIVAL", "title": "Invoice LDS-0298",
            "reason": "The real invoice for D3 arrives after the close and grades the accrual.",
            "build": lambda p: make_pdf(
                p, "Invoice LDS-0298",
                [
                    ("p", "Larkspur Design Studio LLC.   Bill to: Northstar Analytics, Inc.   Date: 8 January 2027"),
                    ("table", [["Description", "Amount (USD)"],
                               ["Deliverable D3 - Visual identity system and usage guidelines", money(LS_D["D3"])]]),
                    ("p", f"Total due: USD {money(LS_D['D3'])}."),
                ],
            ),
        },
    ]  # fmt: skip


def corvid_docs() -> list[Doc]:
    ap_rows = [
        ["Corvid Security - invoice history"],
        ["Invoice", "Date", "Description", "Amount (USD)"],
        ["CSEC-2609", "2026-09-30", "Sentinel Business, September", 3250.00],
        ["CSEC-2610", "2026-10-31", "Sentinel Business, October", 3250.00],
        ["CSEC-2611", "2026-11-30", "Sentinel Business, November", 3250.00],
    ]
    return [
        {
            "n": "01", "name": "corvid_subscription_agreement.docx", "kind": "agreement", "fmt": "DOCX",
            "at": "2026-01-08T09:00:00Z", "role": "SUPPORTS_AMOUNT", "title": "Subscription Agreement",
            "reason": "The subscription agreement's order table sets the original monthly fee and the term.",
            "build": lambda p: make_docx(
                p, "Subscription Agreement - Corvid Security Platform",
                [
                    ("p", "Between Corvid Security, LLC and Northstar Analytics, Inc. Effective 1 January 2026."),
                    ("h", "Order details"),
                    ("table", [["Plan", "Endpoints", "Monthly fee (USD)", "Billing"],
                               ["Sentinel Business", "Up to 250", "3,250.00", "Monthly in arrears"]]),
                    ("h", "Term"),
                    ("p", "The subscription runs from 1 January 2026 to 31 December 2027."),
                    ("h", "Changes"),
                    ("p", "A change to the plan or the fee requires a written amendment agreed by both parties. An email exchange in which both parties confirm the change counts as a written amendment."),
                ],
            ),
        },
        {
            "n": "02", "name": "re_sentinel_enterprise_upgrade.eml", "kind": "email", "fmt": "EML",
            "at": "2026-11-24T16:00:00Z", "role": "SUPPORTS_AMOUNT", "title": "Sentinel Enterprise upgrade",
            "reason": "The confirmed email amendment sets the new monthly fee and the date it takes effect.",
            "build": lambda p: make_eml(p, "FILE-CORVID-02", [
                {"frm": "Avery Rivera <avery.rivera@northstar.example>", "to": "Dana Whitcomb <dana.whitcomb@corvidsec.example>",
                 "at": "2026-11-24T16:00:00Z", "subject": "Re: Sentinel Enterprise upgrade - order confirmation",
                 "body": "Agreed. Please proceed with the upgrade on those terms.\n\nAvery"},
                {"frm": "Dana Whitcomb <dana.whitcomb@corvidsec.example>", "to": "Avery Rivera <avery.rivera@northstar.example>",
                 "at": "2026-11-24T11:30:00Z", "subject": "Sentinel Enterprise upgrade - order confirmation",
                 "body": "Hi Avery, confirming the upgrade from Sentinel Business to Sentinel Enterprise.\nFrom 16 December 2026 the monthly fee is USD 3,900.00, up from USD 3,250.00.\nDecember will be prorated by calendar day, so the invoice will show two rate periods.\nReply with 'agreed' to confirm.\n\nDana Whitcomb, Corvid Security"},
            ]),
        },
        {
            "n": "03", "name": "upgrade_approval_note_corvid.docx", "kind": "brief", "fmt": "DOCX",
            "at": "2026-11-25T09:00:00Z", "role": "SUPPORTS_AMOUNT", "title": "Upgrade approval note",
            "reason": "Internal approval of the upgrade with the effective date and the proration instruction.",
            "build": lambda p: make_docx(
                p, "Upgrade approval note - Corvid Security",
                [
                    ("p", "Approved 25 November 2026 by Avery Rivera (Procurement) and Riley Kim (Engineering service owner)."),
                    ("p", "The plan moves from Sentinel Business to Sentinel Enterprise. The new monthly fee is USD 3,900.00 from 16 December 2026."),
                    ("p", "Finance: accrue December prorated by calendar day, using the old fee up to 15 December and the new fee from 16 December."),
                ],
            ),
        },
        {
            "n": "04", "name": "corvid_order_form_2024.pdf", "kind": "agreement", "fmt": "PDF",
            "at": "2026-12-01T08:00:00Z", "role": "HARD_NEGATIVE", "title": "Order Form 2024-071",
            "reason": "An expired 2024 starter plan at USD 2,900 a month, replaced by the 2026 agreement.",
            "build": lambda p: make_pdf(
                p, "Order Form 2024-071",
                [
                    ("p", "Corvid Security, LLC and Northstar Analytics, Inc. Plan: Sentinel Starter."),
                    ("p", "Monthly fee: USD 2,900.00. Term: 1 January 2024 to 31 December 2025."),
                    ("p", "This order form expired on 31 December 2025 and was replaced by the 2026 Subscription Agreement."),
                ],
            ),
        },
        {
            "n": "05", "name": "invoice_csec-2611_nov.pdf", "kind": "invoice", "fmt": "PDF",
            "at": "2026-12-01T08:00:00Z", "role": "HARD_NEGATIVE", "title": "Invoice CSEC-2611",
            "reason": "The November invoice at the old fee covers a closed period.",
            "build": lambda p: make_pdf(
                p, "Invoice CSEC-2611",
                [
                    ("p", "Corvid Security, LLC.   Bill to: Northstar Analytics, Inc.   Date: 30 November 2026"),
                    ("table", [["Description", "Period", "Amount (USD)"],
                               ["Sentinel Business subscription", "1 Nov - 30 Nov 2026", "3,250.00"]]),
                    ("p", "Total due: USD 3,250.00."),
                ],
            ),
        },
        {
            "n": "06", "name": "corvid_systems_invoice_88213.pdf", "kind": "invoice", "fmt": "PDF",
            "at": "2026-12-09T10:00:00Z", "role": "HARD_NEGATIVE", "title": "Invoice 88213",
            "reason": "An invoice from a different vendor with a similar name, for rack hardware.",
            "build": lambda p: make_pdf(
                p, "Invoice 88213",
                [
                    ("p", "Corvid Systems Inc.   Bill to: Northstar Analytics, Inc.   Date: 9 December 2026"),
                    ("table", [["Description", "Quantity", "Amount (USD)"], ["Rack rails and cable trays", "14", "3,410.00"]]),
                    ("p", "Total due: USD 3,410.00."),
                ],
            ),
        },
        {
            "n": "07", "name": "purchase_order_2026_2103.pdf", "kind": "po", "fmt": "PDF",
            "at": "2026-01-10T09:00:00Z", "role": "HARD_NEGATIVE", "title": "Purchase order PO-2026-2103",
            "reason": "The purchase order still carries the pre-amendment price and was never updated.",
            "build": lambda p: make_pdf(
                p, "Purchase Order PO-2026-2103",
                [
                    ("p", "Vendor: Corvid Security, LLC.   Buyer: Northstar Analytics, Inc.   Cost center CC-ENG-100."),
                    ("table", [["Line", "Description", "Qty", "Unit price (USD)", "Total (USD)"],
                               ["1", "Sentinel Business subscription, monthly", "12", "3,250.00", "39,000.00"]]),
                    ("p", "Approved total: USD 39,000.00. Service period: 1 January 2026 to 31 December 2026."),
                ],
            ),
        },
        {
            "n": "08", "name": "vendor_master_corvid.pdf", "kind": "vendor", "fmt": "PDF",
            "at": "2026-12-01T08:00:00Z", "role": "NOISE", "title": "Supplier record",
            "reason": "Static supplier identity and payment terms carry no amount or period evidence.",
            "build": lambda p: make_pdf(
                p, "Supplier master - Corvid Security",
                vendor_master("VEN-CORVID", "Corvid Security, LLC", "billing@corvidsec.example",
                              "Net 30", "1200 Elm Court, Austin, TX 78701"),
            ),
        },
        {
            "n": "09", "name": "slack_export_secops.txt", "kind": "slack", "fmt": "TXT",
            "at": "2026-12-01T08:00:00Z", "role": "NOISE", "title": "Slack export",
            "reason": "Security operations chatter about alerts with no accounting facts.",
            "build": lambda p: make_txt(p, "secops", [
                ("2026-12-08 09:30", "riley", "corvid flagged three impossible-travel logins overnight, all false positives"),
                ("2026-12-08 09:41", "sam", "closing them out, tuning the geo-fence on the dashboard"),
                ("2026-12-17 14:12", "riley", "new console layout after the upgrade is much easier to read"),
            ]),
        },
        {
            "n": "10", "name": "ap_history_corvid.xlsx", "kind": "ap", "fmt": "XLSX",
            "at": "2026-12-01T08:00:00Z", "role": "HARD_NEGATIVE", "title": "Invoice history",
            "reason": "Past invoices show the old fee and tempt a run-rate estimate.",
            "build": lambda p: make_xlsx(p, {"Invoice history": ap_rows}),
        },
        {
            "n": "L1", "name": "invoice_csec-2612_dec.pdf", "kind": "invoice", "fmt": "PDF",
            "at": "2027-01-04T09:00:00Z", "role": "LATE_ARRIVAL", "title": "Invoice CSEC-2612",
            "reason": "The real December invoice arrives after the close and grades the accrual.",
            "build": lambda p: make_pdf(
                p, "Invoice CSEC-2612",
                [
                    ("p", "Corvid Security, LLC.   Bill to: Northstar Analytics, Inc.   Date: 3 January 2027"),
                    ("table", [["Description", "Days", "Amount (USD)"],
                               ["Sentinel Business at 3,250.00, 1 Dec - 15 Dec 2026", "15 of 31", money(CV_FIRST)],
                               ["Sentinel Enterprise at 3,900.00, 16 Dec - 31 Dec 2026", "16 of 31", money(CV_SECOND)]]),
                    ("p", f"Total due: USD {money(CV_DEC)}."),
                ],
            ),
        },
    ]  # fmt: skip


CASES = [
    ("TERRASTACK", "Terrastack Compute", "terrastack", terrastack_docs),
    ("LARKSPUR", "Larkspur Design Studio", "larkspur", larkspur_docs),
    ("CORVID", "Corvid Security", "corvid", corvid_docs),
]


# Shared company policy, copied once from the demo seed so this seed never moves with it.
COMPANY_CONFIG = json.loads(
    r"""
[
 {
  "config_key": "people",
  "config_value_json": [
   {
    "person_id": "CONTROLLER-001",
    "name": "Rudraksh Awasthi",
    "role": "Controller",
    "email": "rudraksh.awasthi@northstar.example"
   },
   {
    "person_id": "AP-001",
    "name": "Jordan Patel",
    "role": "AP Specialist",
    "email": "jordan.patel@northstar.example"
   },
   {
    "person_id": "PROC-001",
    "name": "Avery Rivera",
    "role": "Procurement Manager",
    "email": "avery.rivera@northstar.example"
   },
   {
    "person_id": "ENG-001",
    "name": "Riley Kim",
    "role": "Engineering Service Owner",
    "email": "riley.kim@northstar.example"
   },
   {
    "person_id": "OPS-001",
    "name": "Taylor Morgan",
    "role": "Operations Owner",
    "email": "taylor.morgan@northstar.example"
   },
   {
    "person_id": "MKT-001",
    "name": "Priya Nair",
    "role": "Marketing Owner",
    "email": "priya.nair@northstar.example"
   }
  ],
  "updated_at": "2026-12-01T08:00:00Z"
 },
 {
  "config_key": "accounting_periods",
  "config_value_json": {
   "2026-09": {
    "status": "CLOSED",
    "phase": "HISTORICAL",
    "period_start": "2026-09-01",
    "period_end": "2026-09-30",
    "close_cutoff": "2026-09-30T23:59:00Z"
   },
   "2026-10": {
    "status": "CLOSED",
    "phase": "HISTORICAL",
    "period_start": "2026-10-01",
    "period_end": "2026-10-31",
    "close_cutoff": "2026-10-31T23:59:00Z"
   },
   "2026-11": {
    "status": "CLOSED",
    "phase": "HISTORICAL",
    "period_start": "2026-11-01",
    "period_end": "2026-11-30",
    "close_cutoff": "2026-11-30T23:59:00Z"
   },
   "2026-12": {
    "status": "OPEN",
    "phase": "LIVE",
    "period_start": "2026-12-01",
    "period_end": "2026-12-31",
    "close_cutoff": "2026-12-31T23:59:00Z"
   },
   "2027-01": {
    "status": "OPEN",
    "phase": "LIVE",
    "period_start": "2027-01-01",
    "period_end": "2027-01-31",
    "close_cutoff": "2027-01-31T23:59:00Z"
   }
  },
  "updated_at": "2026-12-01T08:00:00Z"
 },
 {
  "config_key": "approval_thresholds",
  "config_value_json": {
   "controller_review_above_usd": "25000",
   "mandatory_review_above_usd": "100000",
   "de_minimis_threshold_usd": "5000"
  },
  "updated_at": "2026-12-01T08:00:00Z"
 },
 {
  "config_key": "policy_rules",
  "config_value_json": [
   {
    "rule_id": "POL-01",
    "text": "Never auto-post a case that requires Controller review"
   },
   {
    "rule_id": "POL-02",
    "text": "Never accrue when a matching AP invoice exists"
   },
   {
    "rule_id": "POL-03",
    "text": "Require evidence of service receipt for receipt or usage-based spend"
   },
   {
    "rule_id": "POL-04",
    "text": "Require an open accounting period"
   },
   {
    "rule_id": "POL-05",
    "text": "Require a balanced journal entry"
   },
   {
    "rule_id": "POL-06",
    "text": "Require an active contract or PO, or an approved fallback"
   },
   {
    "rule_id": "POL-07",
    "text": "Capitalize equipment over the threshold and depreciate it over its life"
   },
   {
    "rule_id": "POL-08",
    "text": "Record prepaid services as an asset and expense them over the term"
   }
  ],
  "updated_at": "2026-12-01T08:00:00Z"
 },
 {
  "config_key": "allowed_gl_accounts",
  "config_value_json": [
   {
    "account_code": "610100",
    "name": "SaaS Expense",
    "type": "EXPENSE"
   },
   {
    "account_code": "610200",
    "name": "Cloud & Data Expense",
    "type": "EXPENSE"
   },
   {
    "account_code": "610300",
    "name": "Consulting Expense",
    "type": "EXPENSE"
   },
   {
    "account_code": "610400",
    "name": "Office & Supplies Expense",
    "type": "EXPENSE"
   },
   {
    "account_code": "610500",
    "name": "Marketing & Travel Expense",
    "type": "EXPENSE"
   },
   {
    "account_code": "610600",
    "name": "Depreciation Expense",
    "type": "EXPENSE"
   },
   {
    "account_code": "150100",
    "name": "Prepaid Expenses",
    "type": "ASSET"
   },
   {
    "account_code": "150200",
    "name": "Fixed Assets - Computer Equipment",
    "type": "ASSET"
   },
   {
    "account_code": "150290",
    "name": "Accumulated Depreciation",
    "type": "CONTRA_ASSET"
   },
   {
    "account_code": "210100",
    "name": "Accrued Expenses",
    "type": "LIABILITY"
   },
   {
    "account_code": "210200",
    "name": "Accrued Card Expenses",
    "type": "LIABILITY"
   },
   {
    "account_code": "200100",
    "name": "Accounts Payable",
    "type": "LIABILITY"
   }
  ],
  "updated_at": "2026-12-01T08:00:00Z"
 },
 {
  "config_key": "capitalization_policy",
  "config_value_json": {
   "threshold_usd": "5000",
   "useful_life_months": 36
  },
  "updated_at": "2026-12-01T08:00:00Z"
 },
 {
  "config_key": "simulation_clock",
  "config_value_json": {
   "current_time": "2026-12-01T08:00:00Z"
  },
  "updated_at": "2026-12-01T08:00:00Z"
 },
 {
  "config_key": "company_profile",
  "config_value_json": {
   "name": "Northstar Analytics, Inc.",
   "currency": "USD",
   "close_cycle": "CALENDAR_MONTH",
   "entities": [
    "ENT-001"
   ],
   "accounting_basis": "ACCRUAL"
  },
  "updated_at": "2026-12-01T08:00:00Z"
 }
]
"""
)


# ---- structured company data ---------------------------------------------------------------


def _ap(vendor, key, number, when, received, start, end, amount, po, contract, status, desc):
    return {
        "invoice_id": f"INV-{vendor}-{key}", "vendor_id": f"VEN-{vendor}", "invoice_number": number,
        "invoice_date": when, "received_at": received, "service_start_date": start,
        "service_end_date": end, "amount": str(amount), "currency": "USD", "po_id": po,
        "contract_id": contract, "status": status, "duplicate_flag": False, "credit_flag": False,
        "description": desc, "line_items_json": None, "created_at": received,
        "updated_at": received,
    }  # fmt: skip


def _gl(entry_id, period, posting, vendor, kind, status, desc, lines, created, reverses=None):
    return {
        "gl_entry_id": entry_id, "period": period, "posting_date": posting,
        "vendor_id": f"VEN-{vendor}", "obligation_id": None, "entry_type": kind, "status": status,
        "description": desc,
        "lines_json": [
            {"account_code": a, "debit": d, "credit": c, "description": desc} for a, d, c in lines
        ],
        "source_workpaper_id": None, "reversal_of_gl_entry_id": reverses, "created_at": created,
    }  # fmt: skip


def _ap_gl(vendor, name, key, period, posting, invoice_no, amount, expense, desc):
    text = f"{desc} ({invoice_no})"
    return _gl(
        f"GL-{vendor}-{key}", period, posting, vendor, "AP_INVOICE", "POSTED", text,
        [(expense, str(amount), "0.00"), ("200100", "0.00", str(amount))], f"{posting}T06:30:00Z",
    )  # fmt: skip


def _accrual_pair(vendor, name, key, period, month_end, next_first, amount, expense, next_period):
    accrual_id = f"GL-{vendor}-{key}-ACCRUAL"
    accrual = _gl(
        accrual_id, period, month_end, vendor, "ACCRUAL", "REVERSED",
        f"{name} {key} accrual", [(expense, str(amount), "0.00"), ("210100", "0.00", str(amount))],
        f"{month_end}T18:00:00Z",
    )  # fmt: skip
    reversal = _gl(
        f"{accrual_id}-REV", next_period, next_first, vendor, "ACCRUAL_REVERSAL", "POSTED",
        f"Reverse {name} {key} accrual",
        [("210100", str(amount), "0.00"), (expense, "0.00", str(amount))],
        f"{next_first}T05:30:00Z", reverses=accrual_id,
    )  # fmt: skip
    return [accrual, reversal]


def _evidence(
    rid, vendor, contract, po, start, end, kind, qty, unit, accepted, system, who, status, at
):
    return {
        "service_evidence_id": rid, "vendor_id": f"VEN-{vendor}", "contract_id": contract,
        "po_id": po, "service_start_date": start, "service_end_date": end, "evidence_type": kind,
        "quantity": qty, "unit": unit, "accepted_amount": accepted, "source_system": system,
        "confirmed_by_person_id": who, "confirmation_status": status, "created_at": at,
    }  # fmt: skip


def _po(
    po_id, number, vendor, contract, order_type, cc, owner, total, start, end, account, desc, line
):
    return {
        "po_id": po_id, "po_number": number, "vendor_id": f"VEN-{vendor}", "contract_id": contract,
        "status": "OPEN", "order_type": order_type, "entity_id": "ENT-001", "cost_center": cc,
        "po_owner_id": owner, "approved_total": total, "currency": "USD",
        "service_start_date": start, "service_end_date": end, "gl_account": account,
        "description": desc, "line_items_json": [line],
    }  # fmt: skip


def _contract(row_id, contract, version, vendor, name, status, start, end, model, rate, unit,
              freq, pct, pct_date, service_owner, text):  # fmt: skip
    return {
        "contract_row_id": row_id, "contract_id": contract, "contract_version": version,
        "vendor_id": f"VEN-{vendor}", "contract_name": name, "status": status,
        "effective_start_date": start, "effective_end_date": end, "billing_model": model,
        "base_rate": rate, "rate_unit": unit, "billing_frequency": freq,
        "escalator_percent": pct, "escalator_effective_date": pct_date,
        "service_owner_id": service_owner, "procurement_owner_id": "PROC-001",
        "contract_text": text,
    }  # fmt: skip


def build_static() -> tuple[dict, list[dict], list[dict]]:
    ownership = {
        "controller": "CONTROLLER-001",
        "default_ap_owner": "AP-001",
        "vendors": {
            "VEN-TERRASTACK": {"service_owner_id": "ENG-001", "procurement_owner_id": "PROC-001", "po_owner_id": "ENG-001"},
            "VEN-LARKSPUR": {"service_owner_id": "MKT-001", "procurement_owner_id": "PROC-001", "po_owner_id": "MKT-001"},
            "VEN-CORVID": {"service_owner_id": "ENG-001", "procurement_owner_id": "PROC-001", "po_owner_id": "ENG-001"},
        },
    }  # fmt: skip
    config = [dict(row) for row in COMPANY_CONFIG]
    config.insert(
        1,
        {
            "config_key": "ownership_map",
            "config_value_json": ownership,
            "updated_at": config[0]["updated_at"],
        },
    )

    vendors = [
        {"vendor_id": "VEN-TERRASTACK", "vendor_name": "Terrastack Compute", "vendor_category": "CLOUD", "billing_cadence": "USAGE_BASED", "default_currency": "USD", "billing_contact_email": "billing@terrastack.example", "is_active": True},
        {"vendor_id": "VEN-LARKSPUR", "vendor_name": "Larkspur Design Studio", "vendor_category": "CONSULTING", "billing_cadence": "AD_HOC", "default_currency": "USD", "billing_contact_email": "billing@larkspur.example", "is_active": True},
        {"vendor_id": "VEN-CORVID", "vendor_name": "Corvid Security", "vendor_category": "SAAS", "billing_cadence": "MONTHLY", "default_currency": "USD", "billing_contact_email": "billing@corvidsec.example", "is_active": True},
    ]  # fmt: skip

    contracts = [
        _contract(
            "CON-TERRASTACK-V1", "CON-TERRASTACK", 1, "TERRASTACK",
            "Terrastack GPU Compute Master Services Agreement", "ACTIVE", "2026-01-01", "2027-12-31",
            "USAGE_BASED", "2.50", "COMPUTE_HOUR", "VARIABLE", "8", "2026-10-01", "ENG-001",
            "6.1 Customer pays USD 2.50 per GPU compute-hour. 6.2 From 1 October 2026 the charge rises by 8 percent to USD 2.70 per compute-hour.",
        ),
        _contract(
            "CON-LARKSPUR-V1", "CON-LARKSPUR", 1, "LARKSPUR",
            "Larkspur Statement of Work No. 2 - Brand System Design", "ACTIVE", "2026-10-01",
            "2027-03-31", "MILESTONE_BASED", None, "PROJECT", "VARIABLE", None, None, "MKT-001",
            "Fees are payable per accepted Deliverable, not to exceed USD 96,000.00: D1 18,000.00, D2 22,500.00, D3 27,450.00, D4 28,050.00.",
        ),
        _contract(
            "CON-CORVID-V1", "CON-CORVID", 1, "CORVID", "Corvid Sentinel Business Subscription",
            "SUPERSEDED", "2026-01-01", "2026-12-15", "FIXED_FEE", "3250.00", "MONTH",
            "MONTHLY_IN_ARREARS", None, None, "ENG-001",
            "Sentinel Business plan, up to 250 endpoints, monthly fee USD 3,250.00, billed monthly in arrears.",
        ),
        _contract(
            "CON-CORVID-V2", "CON-CORVID", 2, "CORVID", "Corvid Sentinel Enterprise Subscription (amended)",
            "ACTIVE", "2026-12-16", "2027-12-31", "FIXED_FEE", "3900.00", "MONTH",
            "MONTHLY_IN_ARREARS", None, None, "ENG-001",
            "Upgrade to Sentinel Enterprise. From 16 December 2026 the monthly fee is USD 3,900.00. The transition month is prorated by calendar day.",
        ),
    ]  # fmt: skip

    pos = [
        _po(
            "PO-TERRASTACK-2026", "PO-2026-2101", "TERRASTACK", "CON-TERRASTACK", "BLANKET",
            "CC-ENG-100", "ENG-001", "240000.00", "2026-01-01", "2026-12-31", "610200",
            "Terrastack GPU compute under annual commitment",
            {"po_line_id": "PO-2026-2101-001", "item_category": "BLANKET_LIMIT", "gl_account_code": "610200", "quantity_ordered": None, "unit_price": "2.50", "quantity_received": "0", "quantity_billed": "0", "line_description": "GPU compute-hours, on-demand tier, annual commitment", "service_start_date": "2026-01-01", "service_end_date": "2026-12-31", "receipt_required": False, "useful_life_months": None, "in_service_date": None},
        ),
        _po(
            "PO-LARKSPUR-2026", "PO-2026-2102", "LARKSPUR", "CON-LARKSPUR", "PROJECT",
            "CC-MKT-300", "MKT-001", "96000.00", "2026-10-01", "2027-03-31", "610300",
            "Brand system design project, SOW No. 2, not to exceed USD 96,000",
            {"po_line_id": "PO-2026-2102-001", "item_category": "SERVICE", "gl_account_code": "610300", "quantity_ordered": "1", "unit_price": "96000.00", "quantity_received": "0", "quantity_billed": "0", "line_description": "Brand system design services, fee per accepted deliverable", "service_start_date": "2026-10-01", "service_end_date": "2027-03-31", "receipt_required": True, "useful_life_months": None, "in_service_date": None},
        ),
        _po(
            "PO-CORVID-2026", "PO-2026-2103", "CORVID", "CON-CORVID", "FRAMEWORK", "CC-ENG-100",
            "ENG-001", "39000.00", "2026-01-01", "2026-12-31", "610100",
            "Corvid Sentinel Business subscription, 2026",
            {"po_line_id": "PO-2026-2103-001", "item_category": "SERVICE", "gl_account_code": "610100", "quantity_ordered": "12", "unit_price": "3250.00", "quantity_received": "11", "quantity_billed": "11", "line_description": "Sentinel Business subscription, monthly", "service_start_date": "2026-01-01", "service_end_date": "2026-12-31", "receipt_required": False, "useful_life_months": None, "in_service_date": None},
        ),
    ]  # fmt: skip

    evidence = []
    for period, end in (("2026-09", "30"), ("2026-10", "31"), ("2026-11", "30")):
        evidence.append(
            _evidence(
                f"USE-TERRASTACK-{period}",
                "TERRASTACK",
                "CON-TERRASTACK",
                "PO-TERRASTACK-2026",
                f"{period}-01",
                f"{period}-{end}",
                "SYSTEM_USAGE",
                str(HOURS[period]),
                "COMPUTE_HOUR",
                None,
                "ENGINEERING_PLATFORM",
                None,
                "SYSTEM_VERIFIED",
                f"{period}-{end}T23:00:00Z",
            )
        )
    for key, day, amount in (("2026-10", "27", LS_D["D1"]), ("2026-11", "19", LS_D["D2"])):
        evidence.append(
            _evidence(
                f"USE-LARKSPUR-{key}",
                "LARKSPUR",
                "CON-LARKSPUR",
                "PO-LARKSPUR-2026",
                f"{key}-{day}",
                f"{key}-{day}",
                "MILESTONE_ACCEPTANCE",
                None,
                None,
                str(amount),
                "PROJECT_MANAGEMENT",
                "MKT-001",
                "OWNER_CONFIRMED",
                f"{key}-{day}T15:00:00Z",
            )
        )

    invoices = [
        _ap("TERRASTACK", "2026-09", "TSK-90388", "2026-09-30", "2026-10-01T04:30:00Z", "2026-09-01", "2026-09-30", TS_INVOICE["2026-09"], "PO-TERRASTACK-2026", "CON-TERRASTACK", "PAID", "Terrastack GPU compute, September 2026"),
        _ap("TERRASTACK", "2026-10", "TSK-90455", "2026-10-31", "2026-11-01T04:30:00Z", "2026-10-01", "2026-10-31", TS_INVOICE["2026-10"], "PO-TERRASTACK-2026", "CON-TERRASTACK", "PAID", "Terrastack GPU compute, October 2026"),
        _ap("TERRASTACK", "2026-11", "TSK-90512", "2026-11-30", "2026-12-01T04:30:00Z", "2026-11-01", "2026-11-30", TS_INVOICE["2026-11"], "PO-TERRASTACK-2026", "CON-TERRASTACK", "POSTED", "Terrastack GPU compute, November 2026"),
        _ap("LARKSPUR", "2026-10", "LDS-0231", "2026-10-31", "2026-11-02T09:00:00Z", "2026-10-27", "2026-10-27", LS_D["D1"], "PO-LARKSPUR-2026", "CON-LARKSPUR", "PAID", "Larkspur Deliverable D1, brand audit"),
        _ap("LARKSPUR", "2026-11", "LDS-0247", "2026-11-30", "2026-12-01T05:00:00Z", "2026-11-19", "2026-11-19", LS_NOV_INVOICE, "PO-LARKSPUR-2026", "CON-LARKSPUR", "POSTED", "Larkspur Deliverable D2, brand strategy, with expenses"),
        _ap("CORVID", "2026-09", "CSEC-2609", "2026-09-30", "2026-10-01T06:00:00Z", "2026-09-01", "2026-09-30", CV_OLD, "PO-CORVID-2026", "CON-CORVID", "PAID", "Corvid Sentinel Business subscription, September 2026"),
        _ap("CORVID", "2026-10", "CSEC-2610", "2026-10-31", "2026-11-01T06:00:00Z", "2026-10-01", "2026-10-31", CV_OLD, "PO-CORVID-2026", "CON-CORVID", "PAID", "Corvid Sentinel Business subscription, October 2026"),
        _ap("CORVID", "2026-11", "CSEC-2611", "2026-11-30", "2026-12-01T06:00:00Z", "2026-11-01", "2026-11-30", CV_OLD, "PO-CORVID-2026", "CON-CORVID", "POSTED", "Corvid Sentinel Business subscription, November 2026"),
    ]  # fmt: skip

    gl: list[dict] = []
    gl += [
        _ap_gl("TERRASTACK", "Terrastack", "2026-09", "2026-10", "2026-10-01", "TSK-90388", TS_INVOICE["2026-09"], "610200", "Terrastack GPU compute, September 2026"),
        _ap_gl("TERRASTACK", "Terrastack", "2026-10", "2026-11", "2026-11-01", "TSK-90455", TS_INVOICE["2026-10"], "610200", "Terrastack GPU compute, October 2026"),
        _ap_gl("TERRASTACK", "Terrastack", "2026-11", "2026-12", "2026-12-01", "TSK-90512", TS_INVOICE["2026-11"], "610200", "Terrastack GPU compute, November 2026"),
    ]  # fmt: skip
    for key, month_end, next_first, next_period in (
        ("2026-09", "2026-09-30", "2026-10-01", "2026-10"),
        ("2026-10", "2026-10-31", "2026-11-01", "2026-11"),
        ("2026-11", "2026-11-30", "2026-12-01", "2026-12"),
    ):
        gl += _accrual_pair("TERRASTACK", "Terrastack", key, key, month_end, next_first, TS_ACCRUAL[key], "610200", next_period)  # fmt: skip
    gl += [
        _ap_gl("LARKSPUR", "Larkspur", "2026-10", "2026-11", "2026-11-03", "LDS-0231", LS_D["D1"], "610300", "Larkspur Deliverable D1, brand audit"),
        _ap_gl("LARKSPUR", "Larkspur", "2026-11", "2026-12", "2026-12-01", "LDS-0247", LS_NOV_INVOICE, "610300", "Larkspur Deliverable D2, brand strategy, with expenses"),
    ]  # fmt: skip
    gl += _accrual_pair("LARKSPUR", "Larkspur", "2026-10", "2026-10", "2026-10-31", "2026-11-01", LS_D["D1"], "610300", "2026-11")  # fmt: skip
    gl += _accrual_pair("LARKSPUR", "Larkspur", "2026-11", "2026-11", "2026-11-30", "2026-12-01", LS_D["D2"], "610300", "2026-12")  # fmt: skip
    gl += [
        _ap_gl("CORVID", "Corvid", "2026-09", "2026-10", "2026-10-02", "CSEC-2609", CV_OLD, "610100", "Corvid Sentinel Business subscription, September 2026"),
        _ap_gl("CORVID", "Corvid", "2026-10", "2026-11", "2026-11-02", "CSEC-2610", CV_OLD, "610100", "Corvid Sentinel Business subscription, October 2026"),
        _ap_gl("CORVID", "Corvid", "2026-11", "2026-12", "2026-12-01", "CSEC-2611", CV_OLD, "610100", "Corvid Sentinel Business subscription, November 2026"),
    ]  # fmt: skip

    static = {
        "meta": {
            "generator": "scripts/generate_holdout.py",
            "seed": 7,
            "company": "Northstar Analytics, Inc.",
            "currency": "USD",
            "simulation_start": "2026-12-01T08:00:00Z",
            "historical_periods": ["2026-09", "2026-10", "2026-11"],
            "live_periods": ["2026-12", "2027-01"],
        },
        "company_vendors": vendors,
        "company_contracts": contracts,
        "company_purchase_orders": pos,
        "company_service_evidence": evidence,
        "company_non_po_spend": [],
        "company_ap_invoices": invoices,
        "company_gl_entries": gl,
        "company_config": config,
    }

    events = [
        {"event_id": "EVT-TERRASTACK-2026-12-USAGE-PARTIAL", "available_at": "2026-12-23T18:00:00Z", "operation": "INSERT", "table": "company_service_evidence", "record": _evidence("USE-TERRASTACK-2026-12-PARTIAL", "TERRASTACK", "CON-TERRASTACK", "PO-TERRASTACK-2026", "2026-12-01", "2026-12-21", "SYSTEM_USAGE", str(PARTIAL_DEC_HOURS), "COMPUTE_HOUR", None, "ENGINEERING_PLATFORM", None, "PENDING", "2026-12-23T18:00:00Z")},
        {"event_id": "EVT-LARKSPUR-2026-12-ACCEPTANCE", "available_at": "2026-12-29T16:30:00Z", "operation": "INSERT", "table": "company_service_evidence", "record": _evidence("USE-LARKSPUR-2026-12", "LARKSPUR", "CON-LARKSPUR", "PO-LARKSPUR-2026", "2026-12-14", "2026-12-14", "MILESTONE_ACCEPTANCE", None, None, str(LS_D["D3"]), "PROJECT_MANAGEMENT", "MKT-001", "OWNER_CONFIRMED", "2026-12-29T16:30:00Z")},
        {"event_id": "EVT-CORVID-2027-01-INVOICE", "available_at": "2027-01-04T09:00:00Z", "operation": "INSERT", "table": "company_ap_invoices", "record": _ap("CORVID", "2026-12", "CSEC-2612", "2027-01-03", "2027-01-04T09:00:00Z", "2026-12-01", "2026-12-31", CV_DEC, "PO-CORVID-2026", "CON-CORVID", "IN_QUEUE", "Corvid Sentinel subscription, December 2026, prorated")},
        {"event_id": "EVT-TERRASTACK-2027-01-INVOICE", "available_at": "2027-01-06T09:00:00Z", "operation": "INSERT", "table": "company_ap_invoices", "record": _ap("TERRASTACK", "2026-12", "TSK-90671", "2027-01-05", "2027-01-06T09:00:00Z", "2026-12-01", "2026-12-31", TS_INVOICE["2026-12"], "PO-TERRASTACK-2026", "CON-TERRASTACK", "IN_QUEUE", "Terrastack GPU compute, December 2026")},
        {"event_id": "EVT-LARKSPUR-2027-01-INVOICE", "available_at": "2027-01-09T09:00:00Z", "operation": "INSERT", "table": "company_ap_invoices", "record": _ap("LARKSPUR", "2026-12", "LDS-0298", "2027-01-08", "2027-01-09T09:00:00Z", "2026-12-14", "2026-12-14", LS_D["D3"], "PO-LARKSPUR-2026", "CON-LARKSPUR", "IN_QUEUE", "Larkspur Deliverable D3, visual identity system")},
    ]  # fmt: skip

    outreach = [
        {
            "outreach_key": "TERRASTACK-2026-12-USAGE_CONFIRMATION",
            "available_at": "2027-01-03T10:00:00Z",
            "recipient_role": "SERVICE_OWNER",
            "response_text": "Jordan - the rollup finished. December came to 6,900 GPU compute-hours for 1 to 31 December.",
            "parsed_truth": {"resolved": True, "service_received": True, "quantity": "6900", "unit": "COMPUTE_HOUR"},
            "service_evidence_on_response": _evidence("USE-TERRASTACK-2026-12-CONFIRMED", "TERRASTACK", "CON-TERRASTACK", "PO-TERRASTACK-2026", "2026-12-01", "2026-12-31", "SYSTEM_USAGE", "6900", "COMPUTE_HOUR", None, "ENGINEERING_PLATFORM", "ENG-001", "OWNER_CONFIRMED", "2027-01-03T10:00:00Z"),
        }
    ]  # fmt: skip
    return static, events, outreach


# ---- hidden answer keys ----------------------------------------------------------------------


def historical_truth() -> list[dict[str, Any]]:
    def entry(period, vendor, scenario, invoice, actual, cause, arrival, baseline, status=None):
        return {
            "period": period, "vendor_id": f"VEN-{vendor}", "scenario": scenario,
            "invoice_id": invoice, "expected_actual_amount": str(actual),
            "expected_root_cause": cause, "invoice_arrival_time": arrival,
            "expected_baseline_accrual": None if baseline is None else str(baseline),
            "expected_close_status": status,
        }  # fmt: skip

    return [
        entry("2026-09", "TERRASTACK", "usage_before_uplift", "INV-TERRASTACK-2026-09", TS_INVOICE["2026-09"], None, "2026-10-01T04:30:00Z", TS_ACCRUAL["2026-09"]),
        entry("2026-10", "TERRASTACK", "usage_uplift_missed", "INV-TERRASTACK-2026-10", TS_INVOICE["2026-10"], "MISSED_ESCALATOR", "2026-11-01T04:30:00Z", TS_ACCRUAL["2026-10"]),
        entry("2026-11", "TERRASTACK", "usage_uplift_missed", "INV-TERRASTACK-2026-11", TS_INVOICE["2026-11"], "MISSED_ESCALATOR", "2026-12-01T04:30:00Z", TS_ACCRUAL["2026-11"]),
        entry("2026-12", "TERRASTACK", "usage_partial_outreach", "INV-TERRASTACK-2026-12", TS_INVOICE["2026-12"], "USAGE_VARIANCE", "2027-01-06T09:00:00Z", D(PARTIAL_DEC_HOURS) * BASE, "WAITING"),
        entry("2026-10", "LARKSPUR", "milestone_clean", "INV-LARKSPUR-2026-10", LS_D["D1"], None, "2026-11-02T09:00:00Z", LS_D["D1"]),
        entry("2026-11", "LARKSPUR", "milestone_invoice_over_accepted", "INV-LARKSPUR-2026-11", LS_NOV_INVOICE, "SOURCE_DATA_ERROR", "2026-12-01T05:00:00Z", LS_D["D2"]),
        entry("2026-12", "LARKSPUR", "milestone_needs_review", "INV-LARKSPUR-2026-12", LS_D["D3"], None, "2027-01-09T09:00:00Z", LS_D["D3"], "NEEDS_REVIEW"),
        entry("2026-09", "CORVID", "clean_fixed_history", "INV-CORVID-2026-09", CV_OLD, None, "2026-10-01T06:00:00Z", CV_OLD),
        entry("2026-10", "CORVID", "clean_fixed_history", "INV-CORVID-2026-10", CV_OLD, None, "2026-11-01T06:00:00Z", CV_OLD),
        entry("2026-11", "CORVID", "clean_fixed_history", "INV-CORVID-2026-11", CV_OLD, None, "2026-12-01T06:00:00Z", CV_OLD),
        entry("2026-12", "CORVID", "fixed_fee_mid_month_amendment", "INV-CORVID-2026-12", CV_DEC, None, "2027-01-04T09:00:00Z", CV_DEC, "DONE"),
    ]  # fmt: skip


def holdout_truth() -> dict[str, Any]:
    """What a careful accountant expects at each step. Read only by scripts/run_holdout.py."""
    return {
        "note": "Worked out by hand from the case descriptions in generate_holdout.py.",
        "cases": {
            "CASE-TERRASTACK-2026-12": {
                "vendor_id": "VEN-TERRASTACK",
                "purchase_type": "USAGE_BASED",
                "method": "USAGE_TIMES_RATE",
                "at_close": {"accrual": None, "stage": "AWAITING_OUTREACH"},
                "final": {"accrued": str(TS_INVOICE["2026-12"]), "invoice": str(TS_INVOICE["2026-12"]), "variance": "0.00", "root_cause": None, "stage": "CLOSED"},
                "history": {"2026-09": None, "2026-10": "MISSED_ESCALATOR", "2026-11": "MISSED_ESCALATOR"},
                "rule_expected": True,
                "facts": [
                    {"label": "stepped-up rate 2.70 per hour", "keys": ["UNIT_RATE"], "number": "2.70", "date": None},
                    {"label": "base rate 2.50 per hour", "keys": ["UNIT_RATE"], "number": "2.50", "date": None},
                    {"label": "usage to date 4,690 hours", "keys": ["USAGE_QUANTITY"], "number": "4690", "date": None},
                    {"label": "usage covers only to 2026-12-21", "keys": ["USAGE_COVERAGE_END", "EVIDENCE_GAP"], "number": None, "date": None},
                ],
            },
            "CASE-LARKSPUR-2026-12": {
                "vendor_id": "VEN-LARKSPUR",
                "purchase_type": "MILESTONE_BASED",
                "method": "MILESTONE_ACCEPTED_AMOUNT",
                "at_close": {"accrual": str(LS_D["D3"]), "stage": "AWAITING_ACTUAL_INVOICE", "controller_decision": True},
                "final": {"accrued": str(LS_D["D3"]), "invoice": str(LS_D["D3"]), "variance": "0.00", "root_cause": None, "stage": "CLOSED"},
                "history": {"2026-10": None, "2026-11": "SOURCE_DATA_ERROR"},
                "rule_expected": False,
                "facts": [
                    {"label": "not-to-exceed budget 96,000", "keys": ["BUDGET_CEILING", "ORDER_TOTAL"], "number": "96000", "date": None},
                    {"label": "D3 accepted, fee 27,450", "keys": ["DELIVERED_AMOUNT", "INVOICE_AMOUNT", "ORDER_TOTAL"], "number": "27450", "date": None},
                ],
            },
            "CASE-CORVID-2026-12": {
                "vendor_id": "VEN-CORVID",
                "purchase_type": "FIXED_RECURRING",
                "method": "FIXED_CONTRACT_RATE",
                "at_close": {"accrual": str(CV_DEC), "stage": "AWAITING_ACTUAL_INVOICE"},
                "final": {"accrued": str(CV_DEC), "invoice": str(CV_DEC), "variance": "0.00", "root_cause": None, "stage": "CLOSED"},
                "history": {"2026-09": None, "2026-10": None, "2026-11": None},
                "rule_expected": False,
                "facts": [
                    {"label": "amended monthly fee 3,900", "keys": ["MONTHLY_FEE"], "number": "3900", "date": None},
                    {"label": "original monthly fee 3,250", "keys": ["MONTHLY_FEE"], "number": "3250", "date": None},
                    {"label": "amendment effective 2026-12-16", "keys": ["EFFECTIVE_DATE", "MONTHLY_FEE"], "number": None, "date": "2026-12-16"},
                ],
            },
        },
    }  # fmt: skip


# ---- assemble ---------------------------------------------------------------------------------


def size_label(path: Path, fmt: str) -> str:
    if fmt == "PDF":
        pages = len(PdfReader(str(path)).pages)
        return f"{pages} page{'s' if pages != 1 else ''}"
    if fmt == "XLSX":
        rows = sum(1 for line in read_text(path).splitlines() if not line.startswith("#"))
        return f"{rows} rows"
    if fmt == "EML":
        count = 1 + read_text(path).count(" wrote:")
        return f"{count} message{'s' if count != 1 else ''}"
    if fmt == "TXT":
        count = max(len(path.read_text().splitlines()) - 1, 0)
        return f"{count} messages"
    words = len(read_text(path).split())
    pages = max(1, round(words / 450))
    return f"{pages} page{'s' if pages != 1 else ''}"


def generate(out: Path) -> int:
    """Write the whole held-out seed under `out` and return the number of files."""
    if out.exists():
        shutil.rmtree(out)
    (out / "files").mkdir(parents=True)

    static, events, outreach = build_static()
    StaticCompanyData.model_validate(static)
    for raw in events:
        ScenarioEvent.model_validate(raw)
    for raw in outreach:
        OutreachResponse.model_validate(raw)
    history = historical_truth()
    for raw in history:
        HistoricalTruth.model_validate(raw)

    cases: list[CaseEntry] = []
    files: list[FileEntry] = []
    relevance: dict[str, RelevanceEntry] = {}
    for vendor_key, vendor_name, folder, docs in CASES:
        case_id = f"CASE-{vendor_key}-2026-12"
        cases.append(
            CaseEntry(
                case_id=case_id, vendor_id=f"VEN-{vendor_key}", vendor_name=vendor_name,
                title=f"{vendor_name} December accrual", period="2026-12",
            )
        )  # fmt: skip
        (out / "files" / folder).mkdir()
        for doc in docs():
            file_id = f"FILE-{vendor_key}-{doc['n']}"
            path = out / "files" / folder / doc["name"]
            doc["build"](path)
            text = read_text(path)
            body = " ".join(text.split())[:240]
            files.append(
                FileEntry(
                    file_id=file_id, case_id=case_id, vendor_id=f"VEN-{vendor_key}", name=doc["name"],
                    kind=doc["kind"], format=doc["fmt"], path=f"files/{folder}/{doc['name']}",
                    size_label=size_label(path, doc["fmt"]), available_at=ts(doc["at"]),
                    preview={"card": "memo", "title": doc["title"], "subtitle": vendor_name, "body": body},
                )
            )  # fmt: skip
            relevant = doc["role"] in ("SUPPORTS_AMOUNT", "SUPPORTS_DISCREPANCY", "LATE_ARRIVAL")
            relevance[file_id] = RelevanceEntry(
                file_id=file_id, case_id=case_id, relevant=relevant, role=doc["role"],
                in_universe=doc["role"] != "LATE_ARRIVAL", reason=doc["reason"],
            )  # fmt: skip

    universe = FileUniverse(as_of=ts("2026-12-31T23:59:59Z"), cases=cases, files=files)
    truth = RelevanceTruth(entries=relevance)
    for case in cases:
        in_universe = [e for e in truth.for_case(case.case_id) if e.in_universe]
        wanted = [e for e in in_universe if e.relevant]
        assert len(in_universe) == 10 and 3 <= len(wanted) <= 4, case.case_id

    def dump(name: str, payload: Any) -> None:
        (out / name).write_text(json.dumps(payload, indent=2) + "\n")

    dump(
        "static_company_data.json", StaticCompanyData.model_validate(static).model_dump(mode="json")
    )
    dump(
        "scenario_events.json",
        [ScenarioEvent.model_validate(e).model_dump(mode="json") for e in events],
    )
    dump(
        "outreach_responses.json",
        [OutreachResponse.model_validate(o).model_dump(mode="json") for o in outreach],
    )
    dump("file_universe.json", universe.model_dump(mode="json"))
    dump("relevance_truth.json", truth.model_dump(mode="json"))
    dump(
        "historical_truth.json",
        {
            "expected_outcomes": [
                HistoricalTruth.model_validate(h).model_dump(mode="json") for h in history
            ]
        },
    )
    dump("holdout_truth.json", holdout_truth())
    return len(files)


def main() -> None:
    count = generate(OUT)
    print(f"wrote {count} files for {len(CASES)} cases to {OUT.relative_to(ROOT.parent)}")


if __name__ == "__main__":
    main()
